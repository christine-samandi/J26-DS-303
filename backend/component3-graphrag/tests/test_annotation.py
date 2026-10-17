"""Phase 2: annotation workbook and precision / recall scoring."""

import json

import pytest

pytest.importorskip("spacy")
pytest.importorskip("openpyxl")

from entity_recognition.annotation import (  # noqa: E402
    CORE,
    SHEET,
    build_workbook,
    read_annotations,
    select_chunks,
)
from entity_recognition.chunking import Chunk  # noqa: E402
from entity_recognition.evaluate import Tally, evaluate, report, save  # noqa: E402


def _chunk(chunk_id, text, doc="DOC-IRA-2017"):
    return Chunk(chunk_id, doc, "SEG-X", "section", "1", "1", None, None, 1, None, None, 1, 1, text)


A = "(v) Rs. 1,800,000, for each year of assessment commencing on or after April 1, 2025,"
B = "the personal relief under Section 52 of the principal enactment at 6%"


def _fill(path, annotations, done):
    """Type annotations into the workbook the way the annotator would."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb[SHEET]
    todo = {cid: list(rows) for cid, rows in annotations.items()}
    for row in ws.iter_rows(min_row=2):
        cid, label = row[0].value, row[1].value
        if cid and isinstance(label, str) and label.startswith("CHUNK ") and cid in done:
            row[3].value = "done"
        elif cid and not label and todo.get(cid):
            row[1].value, row[2].value = todo[cid].pop(0)
    wb.save(path)


def test_workbook_round_trip_counts_only_done_chunks(tmp_path):
    path = tmp_path / "ann.xlsx"
    build_workbook([_chunk("CHUNK-A", A), _chunk("CHUNK-B", B)], path)
    _fill(
        path, {"CHUNK-A": [("AMOUNT", "Rs. 1,800,000")], "CHUNK-B": [("RATE", "6%")]}, {"CHUNK-A"}
    )
    texts, gold = read_annotations(path)
    assert texts == {"CHUNK-A": A}
    assert gold == {"CHUNK-A": [("AMOUNT", "Rs. 1,800,000")]}


def test_scores_strict_and_lenient(tmp_path):
    path = tmp_path / "ann.xlsx"
    build_workbook([_chunk("CHUNK-A", A), _chunk("CHUNK-B", B)], path)
    _fill(
        path,
        {
            "CHUNK-A": [("AMOUNT", "Rs. 1,800,000,"), ("DATE", "April 1, 2025")],  # comma ignored
            "CHUNK-B": [
                ("RELIEF", "personal relief"),
                ("SECTION_REF", "Section 52"),  # system finds the longer chain: lenient only
                ("RATE", "6%"),
                ("AMOUNT", "Rs. 99"),  # not in the text: a miss for the system
            ],
        },
        {"CHUNK-A", "CHUNK-B"},
    )
    result = evaluate(*read_annotations(path))
    assert (result.chunks, result.gold) == (2, 6)
    assert result.strict["AMOUNT"].tp == 1 and result.strict["AMOUNT"].fn == 1
    assert result.strict["SECTION_REF"].fp == 1 and result.strict["SECTION_REF"].fn == 1
    assert result.lenient["SECTION_REF"].tp == 1
    assert result.strict["DATE"].tp == result.strict["RATE"].tp == 1
    errors = {(e["error"], e["label"]) for e in result.errors}
    assert ("missed (FN)", "AMOUNT") in errors and ("extra (FP)", "SECTION_REF") in errors
    assert "ALL (micro)" in report(result)
    save(result, tmp_path / "out")
    scores = json.loads((tmp_path / "out" / "extraction_scores.json").read_text())
    assert scores["lenient"]["ALL"]["tp"] == scores["strict"]["ALL"]["tp"] + 1


def test_tally_maths():
    t = Tally(tp=8, fp=2, fn=8)
    assert (t.precision, t.recall) == (0.8, 0.5)
    assert round(t.f1, 4) == round(2 * 0.8 * 0.5 / 1.3, 4)
    assert Tally().f1 == 0.0


def test_sample_is_core_plus_fixed_random():
    chunks = [_chunk(cid, "x" * 200, doc="DOC-IRA-2017") for cid in CORE]
    chunks += [_chunk(f"CHUNK-IRA-2017-S{i:03d}", "y" * 200) for i in range(30)]
    chunks += [_chunk(f"CHUNK-IRA-AMEND-1-S{i:03d}", "z" * 200, doc="DOC-A") for i in range(5)]
    picked = select_chunks(chunks)
    assert [c.chunk_id for c in picked[: len(CORE)]] == CORE
    assert len(picked) == len(CORE) + 8 + 2
    assert [c.chunk_id for c in select_chunks(chunks)] == [c.chunk_id for c in picked]  # seeded


def test_excel_conversions_and_inserted_rows(tmp_path):
    """Excel stores a typed "4%" as 0.04 and "April 1, 2025" as a date; inserted rows may
    have no chunk ID in column A."""
    from datetime import datetime

    from openpyxl import load_workbook

    path = tmp_path / "ann.xlsx"
    build_workbook([_chunk("CHUNK-A", A + " at 4%")], path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    ws["D2"] = "done"
    ws["B4"], ws["C4"] = "RATE", 0.04
    ws["C4"].number_format = "0%"
    ws.insert_rows(5)  # a row the annotator inserted: column A left empty
    ws["B5"], ws["C5"] = "DATE", datetime(2025, 4, 1)
    ws["B6"], ws["C6"] = "AMOUNT", "Rs. 1,800,000"
    wb.save(path)
    texts, gold = read_annotations(path)
    assert gold["CHUNK-A"] == [("RATE", "4%"), ("DATE", "2025-04-01"), ("AMOUNT", "Rs. 1,800,000")]
    result = evaluate(texts, gold)
    assert result.strict["DATE"].tp == result.strict["RATE"].tp == result.strict["AMOUNT"].tp == 1


def test_lenient_prefers_exact_partners():
    from entity_recognition.evaluate import _match

    gold = [("AMOUNT", "Rs. 6"), ("AMOUNT", "Rs. 600,000")]
    pred = [("AMOUNT", "Rs. 600,000"), ("AMOUNT", "Rs. 6")]
    matched, missed, extra = _match(gold, pred, lenient=True)
    assert len(matched) == 2 and not missed and not extra
    assert all(g == p for g, p in matched)
